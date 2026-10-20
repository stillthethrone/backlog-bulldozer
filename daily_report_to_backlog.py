#!/usr/bin/env python3
"""Create a monthly Backlog daily-report issue and post one comment per report date.

The workbook has three sheets:

* ``Backlog Setup``: one row containing the monthly issue fields.
* ``Daily Reports``: one row per historical report date.
* ``Git Activity``: optional source evidence; ignored by the importer.

Use ``--check`` for local validation, omit ``--run`` for an API-backed dry-run,
and add ``--run`` only after reviewing the output.
"""

import argparse
import getpass
import hashlib
import json
import os
import re
import sys
import time
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook


STATE_FILE = Path(__file__).with_name("daily_report_state.json")
BACKUP_FILE = Path(__file__).with_name("daily_report_comment_backup.json")
SETUP_SHEET = "Backlog Setup"
REPORT_SHEET = "Daily Reports"

SETUP_HEADERS = (
    "Project",
    "Summary",
    "Issue Type",
    "Status",
    "Priority",
    "Assignee",
    "Category",
    "Milestone",
    "Version",
    "Start Date",
    "Due Date",
    "Estimated Hours",
    "Actual Hours",
    "Resolution",
    "Existing Issue Key",
)
REPORT_HEADERS = (
    "Report Date",
    "Day",
    "Task",
    "Details",
    "Backlog ID",
    "Duration Hours",
    "Progress",
    "Delay / Issues",
    "Git Repository",
    "Git References",
    "Git URLs",
    "Source",
    "Backlog Comment",
    "Posted Issue Key",
)


def norm(value):
    text = unicodedata.normalize("NFD", str(value or "")).replace("đ", "d").replace("Đ", "D")
    return "".join(char for char in text if unicodedata.category(char) != "Mn").strip().lower()


def one_line(value):
    return " ".join(str(value or "").split())


def as_date(value, label):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        raise ValueError(f"{label} không được để trống")
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"{label} không phải ngày hợp lệ: {value}")


def optional_float(value, label):
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} không phải số: {value}") from exc
    return number


def progress_value(value, label):
    if value in (None, ""):
        raise ValueError(f"{label} không được để trống")
    text = str(value).strip()
    try:
        number = float(text[:-1]) / 100 if text.endswith("%") else float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} không phải phần trăm hợp lệ: {value}") from exc
    if not 0 <= number <= 1:
        raise ValueError(f"{label} phải nằm trong khoảng 0%..100%: {value}")
    return number


def backlog_api_date(value, label, timezone_name):
    """Convert Backlog's date-only or UTC timestamp response to the space date."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        return value
    else:
        text = str(value).strip()
        try:
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
                return date.fromisoformat(text)
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError(f"{label} từ Backlog không phải ngày hợp lệ: {value}") from exc
    if parsed.tzinfo is None:
        return parsed.date()
    try:
        space_timezone = ZoneInfo(timezone_name)
    except (TypeError, ZoneInfoNotFoundError) as exc:
        raise ValueError(f"Timezone của Backlog space không hợp lệ: {timezone_name}") from exc
    return parsed.astimezone(space_timezone).date()


def find_header_row(worksheet, expected, required):
    expected_by_norm = {norm(name): name for name in expected}
    for row in worksheet.iter_rows(min_row=1, max_row=min(30, worksheet.max_row)):
        values = [norm(cell.value) for cell in row]
        if all(norm(name) in values for name in required):
            columns = {
                expected_by_norm[value]: index
                for index, value in enumerate(values)
                if value in expected_by_norm
            }
            return row[0].row, columns
    raise ValueError(f"Không tìm thấy dòng tiêu đề hợp lệ trong sheet '{worksheet.title}'")


def read_workbook(path):
    workbook = load_workbook(path, data_only=True)
    missing_sheets = [name for name in (SETUP_SHEET, REPORT_SHEET) if name not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(f"Thiếu sheet: {', '.join(missing_sheets)}")

    setup_sheet = workbook[SETUP_SHEET]
    setup_row, setup_columns = find_header_row(
        setup_sheet,
        SETUP_HEADERS,
        ("Project", "Summary", "Assignee", "Start Date", "Due Date"),
    )
    setup_values = next(
        setup_sheet.iter_rows(min_row=setup_row + 1, max_row=setup_row + 1, values_only=True),
        None,
    )
    if not setup_values:
        raise ValueError(f"Sheet '{SETUP_SHEET}' chưa có dòng cấu hình")
    setup = {
        header: setup_values[index] if index < len(setup_values) else None
        for header, index in setup_columns.items()
    }
    for header in SETUP_HEADERS:
        setup.setdefault(header, None)
    setup["Start Date"] = as_date(setup["Start Date"], "Start Date")
    setup["Due Date"] = as_date(setup["Due Date"], "Due Date")
    setup["Estimated Hours"] = optional_float(setup["Estimated Hours"], "Estimated Hours")
    setup["Actual Hours"] = optional_float(setup["Actual Hours"], "Actual Hours")

    report_sheet = workbook[REPORT_SHEET]
    report_row, report_columns = find_header_row(
        report_sheet,
        REPORT_HEADERS,
        ("Report Date", "Task", "Details", "Duration Hours", "Progress"),
    )
    reports = []
    for source_row, values in enumerate(
        report_sheet.iter_rows(min_row=report_row + 1, values_only=True),
        start=report_row + 1,
    ):
        date_index = report_columns["Report Date"]
        if date_index >= len(values) or values[date_index] in (None, ""):
            continue
        report = {
            header: values[index] if index < len(values) else None
            for header, index in report_columns.items()
        }
        for header in REPORT_HEADERS:
            report.setdefault(header, None)
        report["Report Date"] = as_date(report["Report Date"], f"Report Date dòng {source_row}")
        report["Duration Hours"] = optional_float(
            report["Duration Hours"],
            f"Duration Hours dòng {source_row}",
        )
        report["Progress"] = progress_value(report["Progress"], f"Progress dòng {source_row}")
        report["_source_row"] = source_row
        reports.append(report)

    validate_workbook(setup, reports)
    return setup, sorted(reports, key=lambda item: item["Report Date"])


def validate_workbook(setup, reports):
    required_setup = ("Project", "Summary", "Issue Type", "Status", "Priority", "Assignee")
    missing = [header for header in required_setup if not one_line(setup.get(header))]
    if missing:
        raise ValueError(f"Thiếu trường cấu hình: {', '.join(missing)}")
    if setup["Start Date"] > setup["Due Date"]:
        raise ValueError("Start Date phải trước hoặc bằng Due Date")
    if not reports:
        raise ValueError(f"Sheet '{REPORT_SHEET}' chưa có báo cáo")

    daily_hours = defaultdict(float)
    for report in reports:
        report_date = report["Report Date"]
        row = report["_source_row"]
        if not setup["Start Date"] <= report_date <= setup["Due Date"]:
            raise ValueError(f"Report Date dòng {row} nằm ngoài Start Date/Due Date")
        for header in ("Task", "Details"):
            if not one_line(report[header]):
                raise ValueError(f"{header} không được để trống tại dòng {row}")
        hours = report["Duration Hours"]
        if hours is None or not 0 < hours <= 24:
            raise ValueError(f"Duration Hours phải nằm trong khoảng (0..24] tại dòng {row}")
        daily_hours[report_date] += hours
    for report_date, total in daily_hours.items():
        if total > 24:
            raise ValueError(f"Tổng Duration Hours ngày {report_date.isoformat()} vượt 24 giờ: {total:g}")


def markdown_list(value, empty_text):
    lines = []
    for raw_line in str(value or "").splitlines():
        line = raw_line.strip().lstrip("-*• ").strip()
        if line:
            lines.append(f"- {line}")
    return "\n".join(lines) if lines else f"- {empty_text}"


def markdown_cell(value):
    return one_line(value).replace("\\", "\\\\").replace("|", "\\|") or "—"


def comment_body(reports):
    if isinstance(reports, dict):
        reports = [reports]
    if not reports:
        raise ValueError("Không có dòng công việc để tạo comment")
    report_date = reports[0]["Report Date"]
    if any(report["Report Date"] != report_date for report in reports):
        raise ValueError("Một comment chỉ được chứa công việc của cùng một ngày")

    explicit_comments = {
        str(report.get("Backlog Comment") or "").strip()
        for report in reports
        if str(report.get("Backlog Comment") or "").strip()
    }
    if len(explicit_comments) > 1:
        raise ValueError(f"Ngày {report_date.isoformat()} có nhiều Backlog Comment override khác nhau")
    if explicit_comments:
        body = explicit_comments.pop()
    else:
        table_rows = []
        finished = []
        issues = []
        ignored_issues = {"", "-", "—", "nothing", "none", "n/a", "na", "khong co", "không có"}
        for report in reports:
            hours = report["Duration Hours"]
            progress = report["Progress"]
            progress_percent = progress * 100
            progress_text = f"{progress_percent:g}%"
            table_rows.append(
                "| "
                + " | ".join((
                    markdown_cell(report["Task"]),
                    markdown_cell(report["Details"]),
                    markdown_cell(report.get("Backlog ID")),
                    f"{hours:g}h",
                    progress_text,
                ))
                + " |"
            )
            if progress == 1:
                finished.append(f"- Finished {one_line(report['Details'])}")
            issue = one_line(report.get("Delay / Issues"))
            if norm(issue) not in ignored_issues and issue not in issues:
                issues.append(issue)

        body = "\n".join((
            f"# {report_date.strftime('%Y/%m/%d')}",
            "",
            "---",
            "",
            "## Daily Report",
            "",
            "---",
            "",
            "| Task | Details | Backlog ID | Duration | Progress |",
            "| --- | --- | --- | ---: | ---: |",
            *table_rows,
            "",
            "## Finish",
            "",
            "---",
            "",
            *(finished or ["- Nothing"]),
            "",
            "## Delay/Issues",
            "",
            "---",
            "",
            *([f"- {issue}" for issue in issues] or ["- Nothing"]),
        ))
    digest = hashlib.sha256(body.encode("utf-8")).hexdigest()
    marker = f"<!-- daily-report:v2:{report_date.isoformat()}:{digest[:12]} -->"
    return f"{body.rstrip()}\n\n{marker}", digest


class Backlog:
    def __init__(self, space, key):
        self.base = f"https://{space}/api/v2"
        self.key = key

    def __call__(self, method, path, **kwargs):
        import requests

        kwargs.setdefault("params", {})["apiKey"] = self.key
        response = requests.request(method, f"{self.base}{path}", timeout=30, **kwargs)
        if not response.ok:
            raise RuntimeError(f"{method} {path} -> {response.status_code}: {response.text}")
        time.sleep(0.25)
        return response.json() if response.text else {}


def exact_item(items, name, label):
    wanted = norm(name)
    matches = [item for item in items if norm(item.get("name")) == wanted]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        available = [item.get("name") for item in items]
        raise ValueError(f"Không tìm thấy {label} '{name}'. Hiện có: {available}")
    raise ValueError(f"{label} '{name}' khớp nhiều giá trị: {[item.get('name') for item in matches]}")


def load_state():
    if not STATE_FILE.exists():
        return {}
    try:
        return json.loads(STATE_FILE.read_text())
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Không đọc được {STATE_FILE.name}: {exc}") from exc


def save_state(state):
    temp = STATE_FILE.with_name(f"{STATE_FILE.name}.tmp")
    temp.write_text(json.dumps(state, ensure_ascii=False, indent=2))
    temp.replace(STATE_FILE)


def state_key(space, project_key, setup, issue_key=None):
    identity = one_line(issue_key or setup["Summary"])
    identity_hash = hashlib.sha256(norm(identity).encode("utf-8")).hexdigest()[:12]
    return ":".join((
        norm(space),
        norm(project_key),
        norm(setup["Assignee"]),
        setup["Start Date"].strftime("%Y-%m"),
        identity_hash,
    ))


MANAGED_MARKER = re.compile(
    r"<!--\s*daily-report:(?:(v\d+):)?(\d{4}-\d{2}-\d{2}):([0-9a-f]{6,64})\s*-->",
    re.IGNORECASE,
)


def managed_comment_records(comments):
    records = []
    for comment in comments:
        content = str(comment.get("content") or "")
        marker = MANAGED_MARKER.search(content)
        if not marker:
            continue
        records.append({
            "id": comment.get("id"),
            "date": marker.group(2),
            "hash": marker.group(3).lower(),
            "version": marker.group(1) or "v1",
            "content": content,
            "created": comment.get("created"),
            "updated": comment.get("updated"),
        })
    return records


def index_existing_comments(comments):
    indexed = {}
    for record in managed_comment_records(comments):
        indexed.setdefault(record["date"], set()).add(record["hash"])

    heading_pattern = re.compile(
        r"(?im)^#{0,6}\s*\[?daily report\]?\s*[-:]?\s*(\d{1,2}/\d{1,2}/\d{4})\s*$"
    )
    for comment in comments:
        content = str(comment.get("content") or "")
        if MANAGED_MARKER.search(content):
            continue
        heading = heading_pattern.search(content)
        if heading:
            try:
                report_date = datetime.strptime(heading.group(1), "%d/%m/%Y").date().isoformat()
            except ValueError:
                continue
            indexed.setdefault(report_date, set()).add("legacy")
    return indexed


def get_all_comments(api, issue_key):
    comments = []
    seen_ids = set()
    cursor = None
    while True:
        params = {"count": 100, "order": "asc"}
        if cursor is not None:
            params["minId"] = cursor
        batch = api("GET", f"/issues/{issue_key}/comments", params=params)
        new_comments = [comment for comment in batch if comment.get("id") not in seen_ids]
        comments.extend(new_comments)
        seen_ids.update(comment.get("id") for comment in new_comments)
        numeric_ids = [comment.get("id") for comment in batch if isinstance(comment.get("id"), int)]
        if len(batch) < 100 or not new_comments or not numeric_ids:
            break
        next_cursor = max(numeric_ids)
        if cursor is not None and next_cursor <= cursor:
            break
        cursor = next_cursor
    return comments


def save_comment_backup(issue_key, comments):
    backup_history = []
    if BACKUP_FILE.exists():
        try:
            backup_history = json.loads(BACKUP_FILE.read_text())
        except (OSError, json.JSONDecodeError) as exc:
            raise ValueError(f"Không đọc được {BACKUP_FILE.name}: {exc}") from exc
        if not isinstance(backup_history, list):
            raise ValueError(f"{BACKUP_FILE.name} phải chứa một danh sách backup")
    backup_history.append({
        "issueKey": issue_key,
        "backedUpAt": datetime.now().astimezone().isoformat(),
        "comments": comments,
    })
    temp = BACKUP_FILE.with_name(f"{BACKUP_FILE.name}.tmp")
    temp.write_text(json.dumps(backup_history, ensure_ascii=False, indent=2))
    temp.replace(BACKUP_FILE)


def build_issue_payload(setup, context):
    payload = {
        "projectId": context["project"]["id"],
        "summary": one_line(setup["Summary"]),
        "description": (
            f"Daily report từ {setup['Start Date'].strftime('%d/%m/%Y')} "
            f"đến {setup['Due Date'].strftime('%d/%m/%Y')}. "
            "Nội dung từng ngày được ghi dưới dạng comment."
        ),
        "issueTypeId": context["issue_type"]["id"],
        "priorityId": context["priority"]["id"],
        "assigneeId": context["assignee"]["id"],
        "startDate": setup["Start Date"].isoformat(),
        "dueDate": setup["Due Date"].isoformat(),
    }
    if context.get("category"):
        payload["categoryId[]"] = [context["category"]["id"]]
    if context.get("milestone"):
        payload["milestoneId[]"] = [context["milestone"]["id"]]
    if context.get("version"):
        payload["versionId[]"] = [context["version"]["id"]]
    if setup["Estimated Hours"] is not None:
        payload["estimatedHours"] = setup["Estimated Hours"]
    if setup["Actual Hours"] is not None:
        payload["actualHours"] = setup["Actual Hours"]
    return payload


def validate_existing_issue(issue, setup, context):
    mismatches = []

    def compare(label, actual, expected):
        if norm(actual) != norm(expected):
            mismatches.append(f"{label}: '{actual}' != '{expected}'")

    compare("Summary", issue.get("summary"), setup["Summary"])
    compare("Issue Type", (issue.get("issueType") or {}).get("name"), setup["Issue Type"])
    compare("Priority", (issue.get("priority") or {}).get("name"), setup["Priority"])
    compare("Assignee", (issue.get("assignee") or {}).get("name"), setup["Assignee"])
    timezone_name = (context.get("space") or {}).get("timezone")
    for label, api_field in (("Start Date", "startDate"), ("Due Date", "dueDate")):
        actual = backlog_api_date(issue.get(api_field), label, timezone_name)
        expected = setup[label]
        if actual != expected:
            mismatches.append(f"{label}: {actual} != {expected}")

    expected_categories = [one_line(setup["Category"])] if one_line(setup["Category"]) else []
    actual_categories = [one_line(item.get("name")) for item in issue.get("category") or []]
    if sorted(map(norm, actual_categories)) != sorted(map(norm, expected_categories)):
        mismatches.append(f"Category: {actual_categories} != {expected_categories}")

    expected_milestones = [one_line(setup["Milestone"])] if one_line(setup["Milestone"]) else []
    actual_milestones = [one_line(item.get("name")) for item in issue.get("milestone") or []]
    if sorted(map(norm, actual_milestones)) != sorted(map(norm, expected_milestones)):
        mismatches.append(f"Milestone: {actual_milestones} != {expected_milestones}")

    expected_versions = [one_line(setup["Version"])] if one_line(setup["Version"]) else []
    actual_versions = [one_line(item.get("name")) for item in issue.get("versions") or []]
    if sorted(map(norm, actual_versions)) != sorted(map(norm, expected_versions)):
        mismatches.append(f"Version: {actual_versions} != {expected_versions}")

    for label, api_field in (("Estimated Hours", "estimatedHours"), ("Actual Hours", "actualHours")):
        expected = setup[label]
        actual = issue.get(api_field)
        if expected is None and actual is None:
            continue
        if expected is None or actual is None or float(actual) != float(expected):
            mismatches.append(f"{label}: {actual} != {expected}")

    expected_resolution = one_line(setup["Resolution"])
    actual_resolution = one_line((issue.get("resolution") or {}).get("name"))
    if norm(actual_resolution) != norm(expected_resolution):
        mismatches.append(f"Resolution: '{actual_resolution}' != '{expected_resolution}'")

    if issue.get("projectId") != context["project"]["id"]:
        mismatches.append(f"Project ID: {issue.get('projectId')} != {context['project']['id']}")
    if mismatches:
        raise ValueError("Issue có sẵn không khớp Backlog Setup:\n- " + "\n- ".join(mismatches))


def resolve_context(api, setup, project_key):
    space = api("GET", "/space")
    project = api("GET", f"/projects/{project_key}")
    issue_type = exact_item(
        api("GET", f"/projects/{project_key}/issueTypes"),
        setup["Issue Type"],
        "Issue Type",
    )
    priority = exact_item(api("GET", "/priorities"), setup["Priority"], "Priority")
    assignee = exact_item(
        api("GET", f"/projects/{project_key}/users"),
        setup["Assignee"],
        "Assignee",
    )
    statuses = api("GET", f"/projects/{project_key}/statuses")
    status = exact_item(statuses, setup["Status"], "Status")
    versions = api("GET", f"/projects/{project_key}/versions")
    context = {
        "space": space,
        "project": project,
        "issue_type": issue_type,
        "priority": priority,
        "assignee": assignee,
        "status": status,
        "category": None,
        "milestone": None,
        "version": None,
    }
    if one_line(setup["Category"]):
        context["category"] = exact_item(
            api("GET", f"/projects/{project_key}/categories"),
            setup["Category"],
            "Category",
        )
    if one_line(setup["Milestone"]):
        context["milestone"] = exact_item(versions, setup["Milestone"], "Milestone")
    if one_line(setup["Version"]):
        context["version"] = exact_item(versions, setup["Version"], "Version")
    return context


def print_local_check(path, setup, reports):
    dates = sorted({report["Report Date"] for report in reports})
    no_git_dates = {
        report["Report Date"]
        for report in reports
        if not one_line(report.get("Git References"))
    }
    total_hours = sum(report["Duration Hours"] for report in reports)
    print(
        f"Đọc {path}: {len(reports)} dòng công việc thuộc {len(dates)} ngày "
        f"từ {dates[0]} đến {dates[-1]}"
    )
    print(f"Issue: {setup['Summary']} | {setup['Issue Type']} | {setup['Status']} | {setup['Assignee']}")
    print(f"Tổng thời lượng: {total_hours:g} giờ | Ngày có dòng chưa gắn Git reference: {len(no_git_dates)}")
    print("File hợp lệ. Chưa gọi Backlog API và chưa tạo issue/comment.")


def main():
    parser = argparse.ArgumentParser(description="Tạo Daily Report trên Backlog từ Excel")
    parser.add_argument("--file", required=True, help="Đường dẫn file Excel Daily Report")
    parser.add_argument("--space", default=os.environ.get("BACKLOG_SPACE", "d-soft-prj.backlog.com"))
    parser.add_argument("--project", default=os.environ.get("BACKLOG_PROJECT"))
    parser.add_argument("--issue", help="Issue key có sẵn, ví dụ HRM-319")
    parser.add_argument("--date", help="Chỉ xử lý một ngày YYYY-MM-DD")
    parser.add_argument(
        "--replace-managed-comments",
        action="store_true",
        help="Xóa comment Daily Report do công cụ quản lý trong phạm vi chọn rồi đăng lại",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--check", action="store_true", help="Chỉ kiểm tra Excel, không gọi API")
    mode.add_argument("--run", action="store_true", help="Tạo issue/comment thật")
    args = parser.parse_args()

    if args.check and args.replace_managed_comments:
        sys.exit("--replace-managed-comments cần dry-run có API hoặc --run; không dùng cùng --check")

    try:
        setup, reports = read_workbook(args.file)
        if args.date:
            selected = as_date(args.date, "--date")
            reports = [report for report in reports if report["Report Date"] == selected]
            if not reports:
                raise ValueError(f"Không tìm thấy Daily Report ngày {selected.isoformat()}")
    except (OSError, ValueError) as exc:
        sys.exit(str(exc))

    if args.check:
        print_local_check(args.file, setup, reports)
        return

    project_key = one_line(args.project or setup["Project"])
    if args.project and norm(args.project) != norm(setup["Project"]):
        sys.exit(f"--project '{args.project}' khác Project trong Excel '{setup['Project']}'")
    issue_key = one_line(args.issue or setup.get("Existing Issue Key")) or None
    key = os.environ.get("BACKLOG_API_KEY") or getpass.getpass("Nhập Backlog API Key: ").strip()
    if not key:
        sys.exit("API Key không được để trống")
    api = Backlog(args.space, key)

    try:
        context = resolve_context(api, setup, project_key)
        payload = build_issue_payload(setup, context)
    except (RuntimeError, ValueError) as exc:
        sys.exit(str(exc))

    try:
        state = load_state()
    except ValueError as exc:
        sys.exit(str(exc))
    month_key = state_key(args.space, project_key, setup, issue_key)
    entry = state.setdefault(month_key, {"comments": {}})
    state_issue = entry.get("issueKey")
    if state_issue and issue_key and state_issue != issue_key:
        sys.exit(f"Issue trong state là {state_issue}, khác với issue được chọn {issue_key}")
    issue_key = issue_key or state_issue

    if issue_key:
        try:
            issue = api("GET", f"/issues/{issue_key}")
        except RuntimeError as exc:
            sys.exit(str(exc))
        try:
            validate_existing_issue(issue, setup, context)
        except ValueError as exc:
            sys.exit(str(exc))
        print(f"Dùng issue có sẵn: {issue_key}")
    else:
        print("Tạo issue:", {key: value for key, value in payload.items() if key != "description"})
        if args.run:
            issue = api("POST", "/issues", data=payload)
            issue_key = issue["issueKey"]
            entry.update(issueId=issue["id"], issueKey=issue_key)
            save_state(state)
            print("  đã tạo:", issue_key)
        else:
            issue_key = "DRY-RUN"

    current_status = norm(issue.get("status", {}).get("name")) if issue_key != "DRY-RUN" else "open"
    desired_status = norm(setup["Status"])
    if current_status != desired_status:
        print(f"Đổi trạng thái: {current_status or '(trống)'} -> {setup['Status']}")
        if args.run:
            api("PATCH", f"/issues/{issue_key}", data={"statusId": context["status"]["id"]})

    comments = []
    existing_markers = {}
    if issue_key != "DRY-RUN":
        comments = get_all_comments(api, issue_key)
        existing_markers = index_existing_comments(comments)

    selected_dates = {report["Report Date"].isoformat() for report in reports}
    if args.replace_managed_comments:
        targets = [
            record
            for record in managed_comment_records(comments)
            if record["date"] in selected_dates
        ]
        print(f"Thay thế comment do công cụ quản lý: {len(targets)} comment trên {issue_key}")
        for target in targets:
            print(f"  xóa {target['date']}: comment {target['id']} ({target['version']})")
        if args.run and targets:
            try:
                save_comment_backup(issue_key, targets)
                print(f"  đã backup vào {BACKUP_FILE.name}")
                for target in targets:
                    api("DELETE", f"/issues/{issue_key}/comments/{target['id']}")
                    entry.setdefault("comments", {}).pop(target["date"], None)
                    save_state(state)
            except (OSError, RuntimeError, ValueError) as exc:
                sys.exit(str(exc))
        target_ids = {target["id"] for target in targets}
        comments = [comment for comment in comments if comment.get("id") not in target_ids]
        existing_markers = index_existing_comments(comments)
        for date_key in selected_dates:
            entry.setdefault("comments", {}).pop(date_key, None)
        if args.run:
            save_state(state)

    reports_by_date = defaultdict(list)
    for report in reports:
        reports_by_date[report["Report Date"]].append(report)

    for report_date in sorted(reports_by_date):
        date_reports = reports_by_date[report_date]
        body, digest = comment_body(date_reports)
        date_key = report_date.isoformat()
        saved = entry.setdefault("comments", {}).get(date_key)
        if saved and saved.get("hash") == digest:
            print(f"  bỏ qua {date_key}: state đã ghi comment {saved.get('commentId')}")
            continue
        remote_hashes = existing_markers.get(date_key, set())
        if digest[:12] in remote_hashes:
            print(f"  bỏ qua {date_key}: comment đã có trên Backlog")
            continue
        if "legacy" in remote_hashes:
            print(f"  bỏ qua {date_key}: đã có comment Daily Report thủ công trên Backlog")
            continue
        if remote_hashes:
            sys.exit(f"Backlog đã có Daily Report {date_key} với nội dung khác; hãy xử lý comment cũ thủ công")
        if saved and saved.get("hash") != digest:
            sys.exit(f"Nội dung {date_key} đã thay đổi sau khi đăng; hãy xử lý comment cũ thủ công")
        print(f"  comment {date_key}: {len(date_reports)} dòng công việc")
        if args.run:
            created = api("POST", f"/issues/{issue_key}/comments", data={"content": body})
            entry["comments"][date_key] = {"hash": digest, "commentId": created["id"]}
            entry.update(issueId=issue.get("id") or entry.get("issueId"), issueKey=issue_key)
            save_state(state)

    print("\nXong." if args.run else "\nĐây là dry-run. Thêm --run để tạo issue/comment thật.")


if __name__ == "__main__":
    main()
