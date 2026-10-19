#!/usr/bin/env python3
"""
Đọc file kế hoạch Sprint -> tạo Parent (UC) + Child (BE/FE) trên Backlog.
Hỗ trợ định dạng cũ (sheet "Backlog") và định dạng Sprint 6 (sheet/cột được tự nhận diện).

Cài đặt:   pip install openpyxl requests
Kiểm tra:  python3 excel_to_backlog.py --check      (chỉ đọc Excel, không gọi API)
Chạy thử:  python3 excel_to_backlog.py              (dry-run: gọi API đọc dữ liệu, không tạo gì)
Chạy thật: python3 excel_to_backlog.py --run

Tuỳ chọn:
  --file PATH        đường dẫn Excel (mặc định ~/Downloads/Sprint_3_Planning_Backlog_UC_lich_du_kien.xlsx)
  --space HOST       mặc định d-soft-prj.backlog.com
  --project KEY      mặc định HRM
  --milestone NAME   mặc định "Sprint 3" (truyền "" để bỏ qua)
  --all              tạo cả task đã Done (mặc định bỏ qua Done)
  --uc UC            chỉ xử lý đúng một UC
  --check            chỉ kiểm tra Excel, không gọi API
API Key được hỏi khi chạy (không hiện trên màn hình), hoặc lấy từ biến BACKLOG_API_KEY.
"""
import argparse, getpass, json, os, sys, time, unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_FILE = Path.home() / "Downloads" / "Sprint_3_Planning_Backlog_UC_lich_du_kien.xlsx"
STATE_FILE = Path(__file__).with_name("created_issues.json")  # chạy lại không tạo trùng
HOURS_PER_DAY = 8
ISSUE_TYPE_NAME = "Task"
PRIORITY_MAP = {"critical": "High", "high": "High", "medium": "Normal", "low": "Low"}  # Backlog chỉ có High/Normal/Low
STATUS_TARGET = {
    "": "open",
    "to do": "open",
    "to-do": "open",
    "open": "open",
    "in progress": "in progress",
    "resolved": "resolved",
    "done": "closed",
    "closed": "closed",
}
HEADER_ALIASES = {
    "id": ("id", "uc / epic"),
    "module": ("module",),
    "kind": ("loai",),
    "method": ("method",),
    "endpoint": ("endpoint / man hinh", "method + endpoint / man hinh"),
    "desc": ("mo ta",),
    "actor": ("actor",),
    "priority": ("priority",),
    "assignee": ("assignee",),
    "status": ("trang thai",),
    "sp": ("est. (sp)", "sp"),
    "hours": ("estimate (gio)", "uoc luong thoi gian (gio)"),
    "days": ("uoc luong thoi gian (ngay)",),
    "start": ("ngay bat dau", "bat dau (du kien)"),
    "end": ("ngay ket thuc", "ket thuc (du kien)", "ket thuc(du kien)"),
}


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()


def one_line(s):
    return " ".join(str(s or "").split())


def target_status(value):
    source = norm(value)
    if source not in STATUS_TARGET:
        raise ValueError(f"Trạng thái Excel chưa được hỗ trợ: '{value}'")
    return STATUS_TARGET[source]


class Backlog:
    def __init__(self, space, key):
        self.base, self.key = f"https://{space}/api/v2", key

    def __call__(self, method, path, **kw):
        import requests

        kw.setdefault("params", {})["apiKey"] = self.key
        r = requests.request(method, f"{self.base}{path}", timeout=30, **kw)
        if not r.ok:
            raise RuntimeError(f"{method} {path} -> {r.status_code}: {r.text}")
        time.sleep(0.25)  # tránh rate limit
        return r.json() if r.text else {}


def read_rows(path):
    wb = load_workbook(path, data_only=True)
    sheets = list(wb.worksheets)
    required = ("id", "module", "kind", "endpoint", "desc", "actor", "priority",
                "assignee", "status", "sp", "start", "end")
    if "Backlog" in wb.sheetnames:
        backlog = wb["Backlog"]
        sheets = [backlog] + [ws for ws in sheets if ws.title != "Backlog"]

    ws = header = hrow = None
    for candidate in sheets:
        for row in candidate.iter_rows(min_row=1, max_row=min(30, candidate.max_row)):
            names = [norm(c.value) for c in row]
            has_schema = all(any(alias in names for alias in HEADER_ALIASES[field]) for field in required)
            has_estimate = any(alias in names for alias in HEADER_ALIASES["hours"] + HEADER_ALIASES["days"])
            if has_schema and has_estimate:
                ws, header, hrow = candidate, names, row[0].row
                break
        if header:
            break

    if not header:
        available = ", ".join(wb.sheetnames)
        sys.exit(f"Không tìm thấy bảng task có cột ID/UC và Assignee. Các sheet: {available}")

    header_idx = {name: i for i, name in enumerate(header) if name}

    def find_column(field):
        return next((header_idx[a] for a in HEADER_ALIASES[field] if a in header_idx), None)

    columns = {field: find_column(field) for field in HEADER_ALIASES}
    missing = [field for field in required if columns[field] is None]
    if columns["hours"] is None and columns["days"] is None:
        missing.append("hours/days")
    if missing:
        sys.exit(f"Thiếu cột trong Excel ({ws.title}): {', '.join(missing)}")

    rows = []
    for source_row, r in enumerate(ws.iter_rows(min_row=hrow + 1, values_only=True), start=hrow + 1):
        if not r[columns["id"]]:
            continue
        d = {
            field: r[columns[field]]
            for field in required
            if field != "endpoint"
        }
        d["id"] = str(d["id"]).strip()
        d["method"] = r[columns["method"]] if columns["method"] is not None else ""
        d["endpoint"] = r[columns["endpoint"]]
        if columns["hours"] is not None:
            raw_estimate = r[columns["hours"]]
            d["hours"] = to_float(raw_estimate)
        else:
            raw_estimate = r[columns["days"]]
            days = to_float(raw_estimate)
            d["hours"] = days * HOURS_PER_DAY if days is not None else None
        if raw_estimate not in (None, "") and d["hours"] is None:
            sys.exit(f"Estimate không phải số tại sheet '{ws.title}', dòng {source_row}: {raw_estimate}")
        rows.append(d)
    return rows


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def fmt_dt(d):
    return d.strftime("%d/%m/%Y %H:%M") if isinstance(d, datetime) else ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", default=str(DEFAULT_FILE))
    ap.add_argument("--space", default=os.environ.get("BACKLOG_SPACE", "d-soft-prj.backlog.com"))
    ap.add_argument("--project", default=os.environ.get("BACKLOG_PROJECT", "HRM"))
    ap.add_argument("--milestone", default="Sprint 3")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--uc", help="Chỉ xử lý đúng một UC, ví dụ UC-12.13")
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--check", action="store_true", help="Chỉ kiểm tra Excel, không gọi Backlog API")
    mode.add_argument("--run", action="store_true", help="Tạo issue thật trên Backlog")
    a = ap.parse_args()

    rows = read_rows(a.file)
    if not a.all:
        rows = [r for r in rows if norm(r["status"]) != "done"]
    if a.uc:
        rows = [r for r in rows if norm(r["id"]) == norm(a.uc)]
        if not rows:
            sys.exit(f"Không tìm thấy task cho UC '{a.uc}'")

    try:
        for row in rows:
            target_status(row["status"])
    except ValueError as exc:
        sys.exit(str(exc))

    mode_name = "CHECK" if a.check else ("tạo thật" if a.run else "DRY-RUN")
    print(f"Đọc {a.file}: {len(rows)} task thuộc {len({r['id'] for r in rows})} UC "
          f"({mode_name})")

    if a.check:
        status_counts = Counter(str(r["status"] or "Open") for r in rows)
        total_hours = sum((r["hours"] or 0) for r in rows)
        status_text = ", ".join(f"{name}: {count}" for name, count in sorted(status_counts.items()))
        print(f"Tổng Estimated Hours: {total_hours:g} giờ")
        print(f"Trạng thái: {status_text}")
        print("File hợp lệ. Chưa gọi Backlog API và chưa tạo issue.")
        return

    key = os.environ.get("BACKLOG_API_KEY") or getpass.getpass("Nhập Backlog API Key: ").strip()
    api = Backlog(a.space, key)

    project = api("GET", f"/projects/{a.project}")
    pid = project["id"]
    types = {t["name"]: t["id"] for t in api("GET", f"/projects/{a.project}/issueTypes")}
    prios = {p["name"]: p["id"] for p in api("GET", "/priorities")}
    users = api("GET", f"/projects/{a.project}/users")
    cats = {c["name"]: c["id"] for c in api("GET", f"/projects/{a.project}/categories")}
    miles = {m["name"]: m["id"] for m in api("GET", f"/projects/{a.project}/versions")}  # "versions" API = Milestone
    statuses = {norm(s["name"]): s["id"] for s in api("GET", f"/projects/{a.project}/statuses")}
    if ISSUE_TYPE_NAME not in types:
        sys.exit(f"Không có Issue Type '{ISSUE_TYPE_NAME}'. Hiện có: {list(types)}")
    required_statuses = {target_status(r["status"]) for r in rows} - {"open"}
    missing_statuses = sorted(name for name in required_statuses if name not in statuses)
    if missing_statuses:
        sys.exit(f"Không tìm thấy trạng thái {missing_statuses} trong project. Hiện có: {list(statuses)}")

    def find_user(name):
        n = norm(name)
        matches = [u for u in users if n and n in norm(u["name"])]
        if len(matches) == 1:
            return matches[0]["id"]
        if len(matches) > 1:
            print(f"  ! Assignee '{name}' khớp nhiều user: {[u['name'] for u in matches]} -> bỏ trống")
            return None
        print(f"  ! Không tìm thấy user '{name}' trong project -> bỏ trống Assignee")

    def ensure(table, path, name):
        if not name:
            return None
        if name not in table:
            if not a.run:
                return None
            table[name] = api("POST", path, data={"name": name})["id"]
        return table[name]

    milestone_id = ensure(miles, f"/projects/{a.project}/versions", a.milestone) if a.milestone else None
    if a.milestone and not milestone_id and not a.run:
        print(f"  (Milestone '{a.milestone}' chưa có, sẽ được tạo khi chạy thật)")

    state = json.loads(STATE_FILE.read_text()) if STATE_FILE.exists() else {}

    def save_state():
        if not a.run:
            return
        temp = STATE_FILE.with_name(f"{STATE_FILE.name}.tmp")
        temp.write_text(json.dumps(state, ensure_ascii=False, indent=2))
        temp.replace(STATE_FILE)

    def create(summary, r, desc, parent_id=None, start=None, end=None, hours=None, cat_names=()):
        data = {
            "projectId": pid, "summary": summary, "description": desc,
            "issueTypeId": types[ISSUE_TYPE_NAME],
            "priorityId": prios[PRIORITY_MAP.get(norm(r["priority"]), "Normal")],
        }
        if parent_id: data["parentIssueId"] = parent_id
        if start: data["startDate"] = start.strftime("%Y-%m-%d")
        if end: data["dueDate"] = end.strftime("%Y-%m-%d")
        if hours is not None: data["estimatedHours"] = hours
        uid = find_user(r["assignee"])
        if uid: data["assigneeId"] = uid
        cids = [c for c in (ensure(cats, f"/projects/{a.project}/categories", n) for n in cat_names) if c]
        if cids: data["categoryId[]"] = cids
        if milestone_id: data["milestoneId[]"] = [milestone_id]
        if not a.run:
            print("    [dry-run]", {k: v for k, v in data.items() if k != "description"})
            return {"id": 0, "issueKey": "DRY-RUN"}
        return api("POST", "/issues", data=data)

    groups = {}
    for r in rows:
        groups.setdefault(r["id"], []).append(r)

    for uc, items in groups.items():
        first = items[0]
        starts = [i["start"] for i in items if isinstance(i["start"], datetime)]
        ends = [i["end"] for i in items if isinstance(i["end"], datetime)]
        total_h = sum((i["hours"] or 0) for i in items) or None
        print(f"\n== {uc} ({len(items)} task)")

        # ---- Parent ----
        st = state.setdefault(uc, {})
        if "id" in st:
            parent_id = st["id"]
            print("  parent đã có:", st["key"])
        else:
            screens = [i["endpoint"] for i in items if i["kind"] == "Frontend"]
            title = one_line(screens[0] if screens else items[0]["endpoint"])
            p = create(f"[{uc}] {first['module']} - {title}", first,
                       desc=f"UC: {uc}\nModule: {first['module']}\nActor: {first['actor'] or ''}\n"
                            f"Số task: {len(items)} ({', '.join(sorted({i['kind'] for i in items}))})",
                       start=min(starts) if starts else None, end=max(ends) if ends else None,
                       hours=total_h, cat_names=[first["module"]])
            parent_id = p["id"]
            st.update(id=p["id"], key=p["issueKey"], children={})
            save_state()
            print("  tạo parent:", p["issueKey"])

        # ---- Children: mỗi dòng Excel = 1 task ----
        for r in items:
            ckey = f"{r['kind']}|{r['method'] or ''}|{r['endpoint']}"
            if ckey in st.get("children", {}):
                print("  child đã có:", ckey)
                continue
            label = {"Backend": "BE", "Frontend": "FE"}.get(r["kind"], r["kind"])
            what = one_line(f"{r['method']} {r['endpoint']}" if r["method"] else r["endpoint"])
            desc = (f"**{r['kind']}** - {what}\n\n{r['desc'] or ''}\n\n"
                    f"- UC: {uc}\n- Module: {r['module']}\n- Actor: {r['actor'] or ''}\n"
                    f"- Priority (Excel): {r['priority']}\n- Est.: {r['sp']} SP\n"
                    f"- Lịch dự kiến: {fmt_dt(r['start'])} -> {fmt_dt(r['end'])}")
            c = create(f"[{uc}][{label}] {what}", r, desc, parent_id=parent_id,
                       start=r["start"] if isinstance(r["start"], datetime) else None,
                       end=r["end"] if isinstance(r["end"], datetime) else None,
                       hours=r["hours"],
                       cat_names=[r["module"], r["kind"]])
            print("  tạo child:", c["issueKey"], f"[{label}] {what}")

            desired_status = target_status(r["status"])
            if desired_status != "open" and a.run:
                sid = statuses.get(desired_status)
                if sid is None:
                    sys.exit(f"Không tìm thấy trạng thái '{desired_status}' trong project. Hiện có: {list(statuses)}")
                upd = {"statusId": sid}
                if desired_status == "closed": upd["resolutionId"] = 0  # 0 = Fixed
                api("PATCH", f"/issues/{c['issueKey']}", data=upd)
            st.setdefault("children", {})[ckey] = c["issueKey"]
            save_state()

        if not a.run:
            state.pop(uc, None)

    print("\nXong." if a.run else "\nĐây là dry-run. Thêm --run để tạo thật.")


if __name__ == "__main__":
    main()
